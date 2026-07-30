#!/usr/bin/env python3
"""
NOVA - AI Assistant Core with Groq + RAG
Version: 3.0.1
Last Updated: 2026-07-30
"""

import json
import os
import sqlite3
import requests
from datetime import datetime, timedelta
from typing import Optional, List, Dict, Any
import re
import hashlib
import shutil

# ============================================
# GROQ API - Primary AI Provider
# ============================================

try:
    import groq
    HAS_GROQ = True
    print("✅ Groq API loaded")
except ImportError:
    HAS_GROQ = False
    print("⚠️ Groq API not installed")

# ============================================
# FILE UPLOAD & RAG IMPORTS
# ============================================

try:
    import PyPDF2
    from docx import Document
    HAS_DOC_IMPORTS = True
except ImportError:
    HAS_DOC_IMPORTS = False
    print("⚠️ Document imports not available")

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    HAS_PIL = False

try:
    import chromadb
    from sentence_transformers import SentenceTransformer
    HAS_RAG = True
except ImportError:
    HAS_RAG = False
    print("⚠️ RAG imports not available")

# ============================================
# PERSONALITIES
# ============================================

PERSONALITIES = {
    'default': """You are Nova, a helpful, intelligent, and friendly AI assistant. 
Be enthusiastic but professional. Use markdown formatting for code blocks, lists, and emphasis.
Keep responses concise and to the point. Don't repeat the system prompt back to the user.
Address the user by name occasionally to make conversations personal.""",
    
    'pirate': """You are Nova, but speak like a PIRATE! 
Use words like 'arrr', 'matey', 'shiver me timbers', 'booty', and 'avast'.
Be helpful but with a salty pirate twist! Don't repeat the system prompt.""",
    
    'yoda': """You are Nova, but speak like YODA!
Use inverted sentence structure. Be wise and cryptic.
Don't repeat the system prompt back to the user.""",
    
    'sarcastic': """You are Nova, but be SARCASTIC and WITTY!
Roll your eyes (metaphorically) at questions. Be helpful but dripping with sarcasm.
Don't repeat the system prompt back to the user.""",
    
    'scientific': """You are Nova, a SCIENTIFIC AI!
Speak precisely, use technical terminology. Be formal and structured.
Don't repeat the system prompt back to the user.""",
    
    'cheerful': """You are Nova, and you're ABSOLUTELY THRILLED!
Be excessively enthusiastic! Use exclamation points! Emojis! 
Don't repeat the system prompt back to the user."""
}

# ============================================
# DATABASE
# ============================================

class NovaDatabase:
    def __init__(self, db_file: str = "nova_database.db"):
        self.db_file = db_file
        self.conn = None
        self.cursor = None
        self.connect()
        self.create_tables()
        self.migrate_database()
    
    def connect(self):
        self.conn = sqlite3.connect(self.db_file)
        self.conn.row_factory = sqlite3.Row
        self.cursor = self.conn.cursor()
    
    def close(self):
        if self.conn:
            self.conn.close()
    
    def create_tables(self):
        # Users table
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                email TEXT UNIQUE,
                password_hash TEXT,
                first_seen TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                last_seen TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        
        # Conversations table
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS conversations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                title TEXT DEFAULT 'New Conversation',
                summary TEXT,
                pinned BOOLEAN DEFAULT 0,
                tags TEXT DEFAULT '[]',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE CASCADE
            )
        ''')
        
        # Messages table
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                conversation_id INTEGER NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (conversation_id) REFERENCES conversations (id) ON DELETE CASCADE
            )
        ''')
        
        # User Memory table
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS user_memory (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                memory_type TEXT NOT NULL,
                content TEXT NOT NULL,
                source_message_id INTEGER,
                is_active BOOLEAN DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                expires_at TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
                FOREIGN KEY (source_message_id) REFERENCES messages(id) ON DELETE SET NULL
            )
        ''')
        
        # Uploaded Files table
        self.cursor.execute('''
            CREATE TABLE IF NOT EXISTS uploaded_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                filename TEXT NOT NULL,
                file_path TEXT NOT NULL,
                file_type TEXT NOT NULL,
                file_size INTEGER,
                content TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
            )
        ''')
        
        self.conn.commit()
    
    def migrate_database(self):
        try:
            self.cursor.execute("PRAGMA table_info(users)")
            columns = [col[1] for col in self.cursor.fetchall()]
            
            if 'email' not in columns:
                self.cursor.execute("ALTER TABLE users ADD COLUMN email TEXT UNIQUE")
                self.conn.commit()
            
            if 'password_hash' not in columns:
                self.cursor.execute("ALTER TABLE users ADD COLUMN password_hash TEXT")
                self.conn.commit()
                
        except Exception as e:
            pass
        
        try:
            self.cursor.execute("PRAGMA table_info(conversations)")
            conv_columns = [col[1] for col in self.cursor.fetchall()]
            
            if 'summary' not in conv_columns:
                self.cursor.execute("ALTER TABLE conversations ADD COLUMN summary TEXT")
                self.conn.commit()
            
            if 'pinned' not in conv_columns:
                self.cursor.execute("ALTER TABLE conversations ADD COLUMN pinned BOOLEAN DEFAULT 0")
                self.conn.commit()
            
            if 'tags' not in conv_columns:
                self.cursor.execute("ALTER TABLE conversations ADD COLUMN tags TEXT DEFAULT '[]'")
                self.conn.commit()
                
        except Exception as e:
            pass
    
    # ============================================
    # USER METHODS
    # ============================================
    
    def get_or_create_user(self, username: str, email: str = None):
        self.cursor.execute("SELECT * FROM users WHERE username = ?", (username,))
        user = self.cursor.fetchone()
        
        if user:
            self.cursor.execute(
                "UPDATE users SET last_seen = CURRENT_TIMESTAMP WHERE id = ?",
                (user['id'],)
            )
            self.conn.commit()
            return dict(user)
        else:
            self.cursor.execute(
                "INSERT INTO users (username, email) VALUES (?, ?)",
                (username, email)
            )
            self.conn.commit()
            user_id = self.cursor.lastrowid
            self.create_conversation(user_id, "First Conversation")
            self.cursor.execute("SELECT * FROM users WHERE id = ?", (user_id,))
            return dict(self.cursor.fetchone())
    
    def get_user_by_email(self, email: str):
        self.cursor.execute("SELECT * FROM users WHERE email = ?", (email,))
        user = self.cursor.fetchone()
        return dict(user) if user else None
    
    def get_user_by_id(self, user_id: int):
        self.cursor.execute("SELECT * FROM users WHERE id = ?", (user_id,))
        user = self.cursor.fetchone()
        return dict(user) if user else None
    
    # ============================================
    # CONVERSATION METHODS
    # ============================================
    
    def create_conversation(self, user_id: int, title: str = "New Conversation"):
        self.cursor.execute(
            "INSERT INTO conversations (user_id, title) VALUES (?, ?)",
            (user_id, title)
        )
        self.conn.commit()
        return self.cursor.lastrowid
    
    def get_conversations(self, user_id: int):
        self.cursor.execute(
            "SELECT * FROM conversations WHERE user_id = ? ORDER BY updated_at DESC",
            (user_id,)
        )
        return [dict(row) for row in self.cursor.fetchall()]
    
    def get_conversation(self, conversation_id: int):
        self.cursor.execute(
            "SELECT * FROM conversations WHERE id = ?",
            (conversation_id,)
        )
        row = self.cursor.fetchone()
        return dict(row) if row else None
    
    def update_conversation_title(self, conversation_id: int, title: str):
        self.cursor.execute(
            "UPDATE conversations SET title = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (title, conversation_id)
        )
        self.conn.commit()
    
    def update_conversation_summary(self, conversation_id: int, summary: str):
        self.cursor.execute(
            "UPDATE conversations SET summary = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (summary, conversation_id)
        )
        self.conn.commit()
    
    def toggle_pin_conversation(self, conversation_id: int, pinned: bool):
        self.cursor.execute(
            "UPDATE conversations SET pinned = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (1 if pinned else 0, conversation_id)
        )
        self.conn.commit()
    
    def delete_conversation(self, conversation_id: int):
        self.cursor.execute("DELETE FROM messages WHERE conversation_id = ?", (conversation_id,))
        self.cursor.execute("DELETE FROM conversations WHERE id = ?", (conversation_id,))
        self.conn.commit()
    
    # ============================================
    # MESSAGE METHODS
    # ============================================
    
    def save_message(self, conversation_id: int, role: str, content: str):
        self.cursor.execute(
            "INSERT INTO messages (conversation_id, role, content) VALUES (?, ?, ?)",
            (conversation_id, role, content)
        )
        message_id = self.cursor.lastrowid
        self.cursor.execute(
            "UPDATE conversations SET updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (conversation_id,)
        )
        self.conn.commit()
        return message_id
    
    def get_messages(self, conversation_id: int, limit: int = 50):
        self.cursor.execute(
            "SELECT * FROM messages WHERE conversation_id = ? ORDER BY timestamp ASC LIMIT ?",
            (conversation_id, limit)
        )
        return [dict(row) for row in self.cursor.fetchall()]
    
    def search_messages(self, user_id: int, query: str):
        self.cursor.execute('''
            SELECT c.title, m.content, m.timestamp 
            FROM messages m
            JOIN conversations c ON m.conversation_id = c.id
            WHERE c.user_id = ? AND m.content LIKE ?
            ORDER BY m.timestamp DESC
        ''', (user_id, f'%{query}%'))
        return [dict(row) for row in self.cursor.fetchall()]
    
    # ============================================
    # FILE METHODS
    # ============================================
    
    def save_uploaded_file(self, user_id: int, filename: str, file_path: str, file_type: str, file_size: int, content: str = None):
        self.cursor.execute('''
            INSERT INTO uploaded_files (user_id, filename, file_path, file_type, file_size, content)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (user_id, filename, file_path, file_type, file_size, content))
        self.conn.commit()
        return self.cursor.lastrowid
    
    def get_uploaded_files(self, user_id: int):
        self.cursor.execute(
            "SELECT * FROM uploaded_files WHERE user_id = ? ORDER BY created_at DESC",
            (user_id,)
        )
        return [dict(row) for row in self.cursor.fetchall()]
    
    def get_uploaded_file(self, file_id: int):
        self.cursor.execute("SELECT * FROM uploaded_files WHERE id = ?", (file_id,))
        row = self.cursor.fetchone()
        return dict(row) if row else None
    
    def delete_uploaded_file(self, file_id: int):
        file = self.get_uploaded_file(file_id)
        if file and os.path.exists(file['file_path']):
            os.remove(file['file_path'])
        self.cursor.execute("DELETE FROM uploaded_files WHERE id = ?", (file_id,))
        self.conn.commit()
    
    # ============================================
    # MEMORY METHODS
    # ============================================
    
    def save_user_memory(self, user_id: int, memory_type: str, content: str, source_message_id: int = None):
        self.cursor.execute('''
            INSERT INTO user_memory (user_id, memory_type, content, source_message_id)
            VALUES (?, ?, ?, ?)
        ''', (user_id, memory_type, content, source_message_id))
        self.conn.commit()
        return self.cursor.lastrowid
    
    def get_user_memories(self, user_id: int, limit: int = 20):
        self.cursor.execute('''
            SELECT * FROM user_memory 
            WHERE user_id = ? AND is_active = 1 AND (expires_at IS NULL OR expires_at > CURRENT_TIMESTAMP)
            ORDER BY updated_at DESC LIMIT ?
        ''', (user_id, limit))
        return [dict(row) for row in self.cursor.fetchall()]
    
    def get_user_memories_by_type(self, user_id: int, memory_type: str):
        self.cursor.execute('''
            SELECT * FROM user_memory 
            WHERE user_id = ? AND memory_type = ? AND is_active = 1
            ORDER BY updated_at DESC
        ''', (user_id, memory_type))
        return [dict(row) for row in self.cursor.fetchall()]
    
    def update_memory(self, memory_id: int, content: str):
        self.cursor.execute('''
            UPDATE user_memory SET content = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?
        ''', (content, memory_id))
        self.conn.commit()
    
    def deactivate_memory(self, memory_id: int):
        self.cursor.execute('''
            UPDATE user_memory SET is_active = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?
        ''', (memory_id,))
        self.conn.commit()
    
    def get_stats(self, user_id: int):
        self.cursor.execute('''
            SELECT COUNT(*) as total_messages FROM messages m
            JOIN conversations c ON m.conversation_id = c.id
            WHERE c.user_id = ?
        ''', (user_id,))
        total_messages = self.cursor.fetchone()['total_messages']
        
        self.cursor.execute(
            "SELECT COUNT(*) as total_conversations FROM conversations WHERE user_id = ?",
            (user_id,)
        )
        total_conversations = self.cursor.fetchone()['total_conversations']
        
        self.cursor.execute(
            "SELECT COUNT(*) as total_memories FROM user_memory WHERE user_id = ? AND is_active = 1",
            (user_id,)
        )
        total_memories = self.cursor.fetchone()['total_memories']
        
        self.cursor.execute(
            "SELECT COUNT(*) as total_files FROM uploaded_files WHERE user_id = ?",
            (user_id,)
        )
        total_files = self.cursor.fetchone()['total_files']
        
        return {
            'total_messages': total_messages,
            'total_conversations': total_conversations,
            'total_memories': total_memories,
            'total_files': total_files
        }

# ============================================
# WEB SEARCH
# ============================================

def web_search(query: str) -> str:
    try:
        from bs4 import BeautifulSoup
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        url = f"https://html.duckduckgo.com/html/?q={query}"
        response = requests.get(url, headers=headers, timeout=10)
        
        if response.status_code == 200:
            soup = BeautifulSoup(response.text, 'html.parser')
            results = soup.find_all('a', class_='result__a')
            snippets = soup.find_all('a', class_='result__snippet')
            
            if results and snippets:
                output = f"🔍 **Search results for '{query}':**\n\n"
                for i in range(min(3, len(results))):
                    title = results[i].get_text(strip=True)
                    snippet = snippets[i].get_text(strip=True) if i < len(snippets) else ""
                    output += f"**{i+1}. {title}**\n{snippet}\n\n"
                return output
            else:
                return f"❌ No results found for '{query}'"
        else:
            return f"❌ Search failed (status: {response.status_code})"
    except Exception as e:
        return f"❌ Search error: {str(e)}"

# ============================================
# CODE EXECUTION
# ============================================

def execute_python_code(code: str) -> str:
    import subprocess
    import sys
    import os
    
    try:
        temp_file = "temp_nova_code.py"
        with open(temp_file, 'w', encoding='utf-8') as f:
            f.write(code)
        
        result = subprocess.run(
            [sys.executable, temp_file],
            capture_output=True,
            text=True,
            timeout=10,
            encoding='utf-8'
        )
        
        os.remove(temp_file)
        
        if result.stdout:
            return f"✅ Output:\n{result.stdout}"
        elif result.stderr:
            return f"❌ Error:\n{result.stderr}"
        else:
            return "✅ Code executed successfully (no output)"
    except subprocess.TimeoutExpired:
        return "❌ Code execution timed out (10s limit)"
    except Exception as e:
        return f"❌ Execution error: {str(e)}"

# ============================================
# FILE PROCESSING
# ============================================

def extract_text_from_file(file_path: str, file_type: str) -> str:
    try:
        if file_type == 'application/pdf':
            if not HAS_DOC_IMPORTS:
                return "PDF support not installed. Run: pip install PyPDF2"
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
        
        elif file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            if not HAS_DOC_IMPORTS:
                return "DOCX support not installed. Run: pip install python-docx"
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        
        elif file_type == 'text/plain':
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        
        elif file_type.startswith('image/'):
            if not HAS_PIL:
                return "Image support not installed. Run: pip install Pillow"
            try:
                image = Image.open(file_path)
                return "Image processed. Content extraction not available."
            except Exception as e:
                return f"Image error: {str(e)}"
        
        else:
            return f"Unsupported file type: {file_type}"
    
    except Exception as e:
        return f"Error extracting text: {str(e)}"

# ============================================
# RAG SYSTEM
# ============================================

class RAGSystem:
    def __init__(self, user_id: int):
        self.user_id = user_id
        self.collection_name = f"user_{user_id}_docs"
        self.chroma_client = None
        self.embedding_model = None
        self.collection = None
        
        if HAS_RAG:
            self._initialize()
        else:
            print("⚠️ RAG not available")
    
    def _initialize(self):
        try:
            self.chroma_client = chromadb.PersistentClient(path="./chroma_db")
            self.embedding_model = SentenceTransformer('all-MiniLM-L6-v2')
            
            try:
                self.collection = self.chroma_client.get_collection(self.collection_name)
            except:
                self.collection = self.chroma_client.create_collection(self.collection_name)
                
        except Exception as e:
            print(f"RAG initialization error: {e}")
    
    def _split_text(self, text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
        chunks = []
        start = 0
        text_length = len(text)
        
        while start < text_length:
            end = start + chunk_size
            if end < text_length:
                while end > start and text[end] not in [' ', '\n', '.', '!', '?']:
                    end -= 1
            chunks.append(text[start:end].strip())
            start = end - overlap
        
        return [chunk for chunk in chunks if len(chunk) > 20]
    
    def add_document(self, file_id: int, content: str, metadata: dict = None):
        if not HAS_RAG or self.collection is None:
            return
        
        try:
            chunks = self._split_text(content)
            
            for i, chunk in enumerate(chunks):
                embedding = self.embedding_model.encode(chunk).tolist()
                self.collection.add(
                    embeddings=[embedding],
                    documents=[chunk],
                    metadatas=[metadata or {}],
                    ids=[f"file_{file_id}_chunk_{i}"]
                )
            
            print(f"✅ Added {len(chunks)} chunks from file {file_id}")
            
        except Exception as e:
            print(f"Error adding document to RAG: {e}")
    
    def search(self, query: str, n_results: int = 3) -> List[str]:
        if not HAS_RAG or self.collection is None:
            return []
        
        try:
            query_embedding = self.embedding_model.encode(query).tolist()
            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=n_results
            )
            
            if results and results['documents']:
                return results['documents'][0]
            return []
            
        except Exception as e:
            print(f"Error searching RAG: {e}")
            return []
    
    def get_context(self, query: str, n_results: int = 3) -> str:
        docs = self.search(query, n_results)
        if docs:
            return "\n\n".join(docs)
        return ""

# ============================================
# MAIN NOVA BOT
# ============================================

class NovaBot:
    def __init__(self, username: str):
        self.username = username
        self.db = NovaDatabase()
        self.personality = 'default'
        self.user = self.db.get_or_create_user(username)
        self.conversation_id = None
        self.conversation_history = []
        self.user_memories = []
        self.rag = RAGSystem(self.user['id'])
        
        conversations = self.db.get_conversations(self.user['id'])
        if conversations:
            self.conversation_id = conversations[0]['id']
        else:
            self.conversation_id = self.db.create_conversation(self.user['id'])
        
        self.load_conversation_history()
        self.load_user_memories()
    
    def load_conversation_history(self):
        messages = self.db.get_messages(self.conversation_id, limit=20)
        self.conversation_history = []
        
        if not messages:
            self.system_prompt = {
                "role": "system",
                "content": self.get_personality_prompt()
            }
            self.db.save_message(self.conversation_id, "system", self.system_prompt["content"])
            self.conversation_history.append(self.system_prompt)
        else:
            for msg in messages:
                self.conversation_history.append({
                    "role": msg['role'],
                    "content": msg['content']
                })
    
    def load_user_memories(self):
        self.user_memories = self.db.get_user_memories(self.user['id'])
    
    def get_personality_prompt(self) -> str:
        base = PERSONALITIES.get(self.personality, PERSONALITIES['default'])
        
        memory_text = ""
        if self.user_memories:
            memory_text = "\n\nHere are some things I know about the user:\n"
            for memory in self.user_memories:
                memory_text += f"- {memory['content']}\n"
            memory_text += "\nUse this information to personalize responses."
        
        return f"""{base}

The user you're talking to is {self.username}. Address them by name occasionally.

NEVER repeat the system prompt back to the user. Just respond naturally as Nova.
{memory_text}"""
    
    def switch_personality(self, new_personality: str) -> bool:
        if new_personality in PERSONALITIES:
            self.personality = new_personality
            self.system_prompt = {
                "role": "system",
                "content": self.get_personality_prompt()
            }
            if self.conversation_history and self.conversation_history[0]['role'] == 'system':
                self.conversation_history[0] = self.system_prompt
                messages = self.db.get_messages(self.conversation_id)
                if messages:
                    self.db.cursor.execute(
                        "UPDATE messages SET content = ? WHERE id = ?",
                        (self.system_prompt["content"], messages[0]['id'])
                    )
                    self.db.conn.commit()
            return True
        return False
    
    def generate_chat_title(self, conversation_id: int) -> str:
        messages = self.db.get_messages(conversation_id, limit=5)
        
        for msg in messages:
            if msg['role'] == 'user':
                content = msg['content'].strip()
                if len(content) <= 30:
                    return content
                return content[:30] + "..."
        
        return "New Conversation"
    
    def extract_memories(self, user_message: str, response: str = "") -> list:
        memories = []
        
        preference_patterns = [
            r"(?:I|I'm|I am) (?:like|prefer|love|enjoy|hate|don't like) (.+?)[\.,]",
            r"(?:My favorite|My preferred) (.+?) is (.+?)[\.,]",
        ]
        
        for pattern in preference_patterns:
            matches = re.findall(pattern, user_message, re.IGNORECASE)
            for match in matches:
                memories.append({
                    'type': 'preference',
                    'content': match if isinstance(match, str) else ' '.join(match)
                })
        
        fact_patterns = [
            r"(?:I|I'm|I am|I have|I've) (.+?)[\.,]",
            r"(?:My|Mine|My name is) (.+?)[\.,]",
        ]
        
        for pattern in fact_patterns:
            matches = re.findall(pattern, user_message, re.IGNORECASE)
            for match in matches:
                if len(match) > 5 and len(match) < 100:
                    memories.append({
                        'type': 'fact',
                        'content': match if isinstance(match, str) else ' '.join(match)
                    })
        
        project_patterns = [
            r"(?:I|I'm|I am) (?:working on|building|creating|developing) (.+?)[\.,]",
            r"(?:My|My current) project (.+?)[\.,]",
        ]
        
        for pattern in project_patterns:
            matches = re.findall(pattern, user_message, re.IGNORECASE)
            for match in matches:
                memories.append({
                    'type': 'project',
                    'content': match if isinstance(match, str) else ' '.join(match)
                })
        
        return memories
    
    def save_memories(self, user_message: str, message_id: int):
        extracted = self.extract_memories(user_message)
        
        for memory in extracted:
            existing = self.db.cursor.execute('''
                SELECT id FROM user_memory 
                WHERE user_id = ? AND memory_type = ? AND content LIKE ? AND is_active = 1
            ''', (self.user['id'], memory['type'], f"%{memory['content']}%"))
            
            if not existing.fetchone():
                self.db.save_user_memory(
                    self.user['id'],
                    memory['type'],
                    memory['content'],
                    message_id
                )
                print(f"🧠 Saved memory: {memory['type']} - {memory['content']}")
        
        self.load_user_memories()
    
    def upload_file(self, filename: str, file_path: str, file_type: str, file_size: int) -> dict:
        try:
            content = extract_text_from_file(file_path, file_type)
            
            file_id = self.db.save_uploaded_file(
                self.user['id'],
                filename,
                file_path,
                file_type,
                file_size,
                content
            )
            
            if content and len(content) > 50 and HAS_RAG:
                metadata = {
                    "filename": filename,
                    "file_type": file_type,
                    "user_id": self.user['id']
                }
                self.rag.add_document(file_id, content, metadata)
            
            return {
                "success": True,
                "file_id": file_id,
                "filename": filename,
                "content_preview": content[:500] if content else "No text extracted"
            }
            
        except Exception as e:
            return {
                "success": False,
                "error": str(e)
            }
    
    def get_response(self, user_input: str) -> str:
        """Get response from AI using Groq API ONLY - NO OLLAMA"""
        
        # Check for special commands
        if user_input.lower().startswith('/search'):
            query = user_input[8:].strip()
            if query:
                return web_search(query)
            else:
                return "Please provide a search query. Example: /search python tutorials"
        
        elif user_input.lower().startswith('/run'):
            code = user_input[4:].strip()
            if code:
                return execute_python_code(code)
            else:
                return "Please provide code to execute. Example: /run print('Hello')"
        
        # Save user message
        message_id = self.db.save_message(self.conversation_id, "user", user_input)
        self.conversation_history.append({"role": "user", "content": user_input})
        
        # Extract and save memories
        self.save_memories(user_input, message_id)
        
        # Build messages for Groq
        messages_to_send = self.conversation_history.copy()
        
        # Add memories
        if self.user_memories:
            memory_text = "Here are some facts about the user: " + \
                          ", ".join([m['content'] for m in self.user_memories[:5]])
            if messages_to_send and messages_to_send[0]['role'] == 'system':
                messages_to_send[0]['content'] += f"\n\nUser information: {memory_text}"
        
        # Add RAG context
        if HAS_RAG and self.rag:
            try:
                docs = self.rag.search(user_input, n_results=2)
                if docs:
                    rag_context = "\n\nHere is some relevant information from uploaded documents:\n" + "\n---\n".join(docs)
                    messages_to_send.insert(1, {
                        "role": "system",
                        "content": rag_context
                    })
            except Exception as e:
                print(f"RAG search error: {e}")
        
        try:
            if not HAS_GROQ:
                return "Error: Groq API not installed. Please install groq: pip install groq"
            
            api_key = os.environ.get("GROQ_API_KEY")
            if not api_key:
                return "Error: GROQ_API_KEY not set in environment variables."
            
            client = groq.Groq(api_key=api_key)
            response = client.chat.completions.create(
                messages=messages_to_send,
                model="llama-3.3-70b-versatile",
                temperature=0.7,
                max_tokens=512,
            )
            full_response = response.choices[0].message.content
            
            # Save assistant response
            self.db.save_message(self.conversation_id, "assistant", full_response)
            self.conversation_history.append({"role": "assistant", "content": full_response})
            
            # Update conversation title if needed
            conversations = self.db.get_conversations(self.user['id'])
            if len(conversations) == 1 and conversations[0]['title'] == "First Conversation":
                title = self.generate_chat_title(self.conversation_id)
                self.db.update_conversation_title(self.conversation_id, title)
            
            return full_response
            
        except Exception as e:
            print(f"❌ Groq API error: {e}")
            return f"Error: {str(e)}. Please check your GROQ_API_KEY environment variable."
    
    def switch_conversation(self, conversation_id: int):
        self.conversation_id = conversation_id
        self.load_conversation_history()
        self.load_user_memories()
    
    def create_new_conversation(self, title: str = "New Conversation") -> int:
        conv_id = self.db.create_conversation(self.user['id'], title)
        self.switch_conversation(conv_id)
        return conv_id
    
    def delete_conversation(self, conversation_id: int):
        if conversation_id == self.conversation_id:
            conversations = self.db.get_conversations(self.user['id'])
            other = [c for c in conversations if c['id'] != conversation_id]
            if other:
                self.switch_conversation(other[0]['id'])
            else:
                self.create_new_conversation()
        
        self.db.delete_conversation(conversation_id)
    
    def get_conversations(self) -> list:
        return self.db.get_conversations(self.user['id'])
    
    def get_messages(self) -> list:
        return self.db.get_messages(self.conversation_id)
    
    def export_conversation(self, format_type: str = 'txt') -> str:
        import os
        from datetime import datetime
        
        messages = self.get_messages()
        if not messages:
            return "No messages to export"
        
        conversations = self.get_conversations()
        title = next((c['title'] for c in conversations if c['id'] == self.conversation_id), "conversation")
        clean_title = re.sub(r'[^\w\s-]', '', title)[:30]
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"nova_exports/{clean_title}_{timestamp}.{format_type}"
        
        if format_type == 'txt':
            export_to_txt(messages, filename)
        elif format_type == 'json':
            export_to_json(messages, filename)
        elif format_type == 'pdf':
            export_to_pdf(messages, filename)
        else:
            return f"Unknown format: {format_type}"
        
        return f"Exported to: {filename}"
    
    def get_stats(self) -> dict:
        return self.db.get_stats(self.user['id'])

# ============================================
# EXPORT FUNCTIONS
# ============================================

def export_to_txt(messages: list, filename: str) -> str:
    import os
    os.makedirs(os.path.dirname(filename) or '.', exist_ok=True)
    
    with open(filename, 'w', encoding='utf-8') as f:
        f.write("=" * 60 + "\n")
        f.write(f"NOVA CONVERSATION EXPORT\n")
        f.write(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write("=" * 60 + "\n\n")
        
        for msg in messages:
            if msg['role'] == "system":
                continue
            role_label = "👤 User" if msg['role'] == "user" else "🤖 Nova"
            f.write(f"[{msg['timestamp']}] {role_label}:\n")
            f.write(f"{msg['content']}\n")
            f.write("-" * 40 + "\n")
    
    return filename

def export_to_json(messages: list, filename: str) -> str:
    import os
    os.makedirs(os.path.dirname(filename) or '.', exist_ok=True)
    
    export_data = {
        'export_date': datetime.now().isoformat(),
        'messages': [
            {
                'role': msg['role'],
                'content': msg['content'],
                'timestamp': msg['timestamp']
            }
            for msg in messages
            if msg['role'] != 'system'
        ]
    }
    
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(export_data, f, indent=2, ensure_ascii=False)
    
    return filename

def export_to_pdf(messages: list, filename: str) -> str:
    from fpdf import FPDF
    import os
    
    os.makedirs(os.path.dirname(filename) or '.', exist_ok=True)
    
    class PDF(FPDF):
        def __init__(self):
            super().__init__()
            self.set_font('Arial', '', 10)
    
    pdf = PDF()
    pdf.add_page()
    
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(0, 10, "NOVA Conversation Export", ln=True, align='C')
    pdf.set_font('Arial', '', 10)
    pdf.cell(0, 10, f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True, align='C')
    pdf.ln(10)
    
    pdf.set_font('Arial', '', 10)
    for msg in messages:
        if msg['role'] == "system":
            continue
        role_label = "User" if msg['role'] == "user" else "Nova"
        pdf.set_font('Arial', 'B', 10)
        pdf.cell(0, 6, f"[{msg['timestamp']}] {role_label}:", ln=True)
        pdf.set_font('Arial', '', 10)
        clean_content = msg['content'].encode('utf-8', 'ignore').decode('utf-8')
        for line in clean_content.split('\n'):
            if line.strip():
                pdf.multi_cell(0, 6, line)
        pdf.ln(4)
    
    pdf.output(filename)
    return filename

# ============================================
# TERMINAL UI
# ============================================

if __name__ == "__main__":
    import sys
    
    try:
        from rich.console import Console
        from rich.panel import Panel
        from rich.table import Table
        HAS_RICH = True
    except ImportError:
        HAS_RICH = False
    
    if HAS_RICH:
        console = Console()
    
    print("\n🚀 Welcome to Nova AI with Groq + RAG!")
    print("📅 Version: 3.0.1 - 2026-07-30")
    username = input("What's your name? ").strip() or "User"
    
    bot = NovaBot(username)
    
    print(f"\n👋 Welcome back, {username}!")
    print(f"💬 Current conversation: {bot.conversation_id}")
    print(f"🧠 Personality: {bot.personality}")
    print(f"📚 Memories: {len(bot.user_memories)} facts stored")
    print("\nCommands:")
    print("  /search <query>  - Search the web")
    print("  /run <code>      - Execute Python code")
    print("  /new             - New conversation")
    print("  /list            - List conversations")
    print("  /switch <id>     - Switch conversation")
    print("  /delete <id>     - Delete conversation")
    print("  /export <format> - Export conversation (txt, json, pdf)")
    print("  /stats           - Show statistics")
    print("  /memories        - Show stored memories")
    print("  /files           - Show uploaded files")
    print("  /exit            - Exit")
    print()
    
    while True:
        try:
            user_input = input(f"\n[You] ").strip()
            
            if not user_input:
                continue
            
            if user_input.lower() in ['/exit', '/quit']:
                print("👋 Goodbye!")
                break
            
            elif user_input == '/new':
                conv_id = bot.create_new_conversation()
                print(f"✅ New conversation created: {conv_id}")
                continue
            
            elif user_input == '/list':
                conversations = bot.get_conversations()
                if HAS_RICH:
                    table = Table(title="📚 Your Conversations")
                    table.add_column("ID", style="cyan")
                    table.add_column("Title", style="white")
                    table.add_column("Updated", style="dim")
                    table.add_column("Pinned", style="dim")
                    for conv in conversations:
                        table.add_row(
                            str(conv['id']),
                            conv['title'],
                            conv['updated_at'][:16],
                            "📌" if conv.get('pinned', 0) else ""
                        )
                    console.print(table)
                else:
                    for conv in conversations:
                        print(f"  {conv['id']}: {conv['title']} ({conv['updated_at'][:16]})")
                continue
            
            elif user_input == '/memories':
                memories = bot.user_memories
                if memories:
                    print(f"\n🧠 Stored memories for {username}:")
                    for i, mem in enumerate(memories, 1):
                        print(f"  {i}. [{mem['memory_type']}] {mem['content']}")
                        print(f"     (Updated: {mem['updated_at'][:16]})")
                else:
                    print("No memories stored yet. Start a conversation to build memories!")
                continue
            
            elif user_input == '/files':
                files = bot.db.get_uploaded_files(bot.user['id'])
                if files:
                    print(f"\n📁 Uploaded files for {username}:")
                    for i, file in enumerate(files, 1):
                        size_kb = file['file_size'] / 1024 if file['file_size'] else 0
                        print(f"  {i}. {file['filename']} ({size_kb:.1f} KB)")
                        print(f"     Type: {file['file_type']}")
                        print(f"     Uploaded: {file['created_at'][:16]}")
                else:
                    print("No files uploaded yet.")
                continue
            
            elif user_input.startswith('/switch '):
                try:
                    conv_id = int(user_input.split()[1])
                    bot.switch_conversation(conv_id)
                    print(f"✅ Switched to conversation: {conv_id}")
                except:
                    print("❌ Invalid conversation ID")
                continue
            
            elif user_input.startswith('/delete '):
                try:
                    conv_id = int(user_input.split()[1])
                    bot.delete_conversation(conv_id)
                    print(f"✅ Deleted conversation: {conv_id}")
                except:
                    print("❌ Invalid conversation ID")
                continue
            
            elif user_input.startswith('/export '):
                format_type = user_input.split()[1] if len(user_input.split()) > 1 else 'txt'
                result = bot.export_conversation(format_type)
                print(f"✅ {result}")
                continue
            
            elif user_input == '/stats':
                stats = bot.get_stats()
                print(f"📊 Statistics for {username}:")
                print(f"  Total Messages: {stats['total_messages']}")
                print(f"  Total Conversations: {stats['total_conversations']}")
                print(f"  Total Memories: {stats['total_memories']}")
                print(f"  Total Files: {stats['total_files']}")
                continue
            
            response = bot.get_response(user_input)
            print(f"\n[Nova] {response}")
            
        except KeyboardInterrupt:
            print("\n👋 Goodbye!")
            break
        except Exception as e:
            print(f"❌ Error: {e}")
    
    bot.db.close()